import pymupdf
import pymupdf4llm
from docx import Document

from backend import document_parser


def test_txt_parser(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text(
        "First line\nSecond line\n",
        encoding="utf-8",
    )

    result = document_parser.parse_document(path)

    assert result.file_type == "txt"
    assert result.extraction_method == "text_reader"
    assert "First line" in result.content
    assert "Second line" in result.content


def test_docx_parser_preserves_heading_paragraph_and_table(tmp_path):
    path = tmp_path / "report.docx"

    document = Document()
    document.add_heading("Security Report", level=1)
    document.add_paragraph("Summary of the findings.")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Finding"
    table.cell(0, 1).text = "Severity"
    table.cell(1, 0).text = "SQL injection"
    table.cell(1, 1).text = "High"

    document.save(path)

    result = document_parser.parse_document(path)

    assert result.extraction_method == "python-docx"
    assert "# Security Report" in result.content
    assert "Summary of the findings." in result.content
    assert "SQL injection" in result.content
    assert "High" in result.content


def test_pdf_parser_uses_pymupdf4llm(tmp_path, monkeypatch):
    path = tmp_path / "sample.pdf"

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "PDF parser test")
    pdf.save(path)
    pdf.close()

    def fake_to_markdown(file_path, *, page_chunks):
        assert page_chunks is True
        return [
            {
                "metadata": {"page": 1},
                "text": "# PDF parser test",
            }
        ]

    monkeypatch.setattr(
        pymupdf4llm,
        "to_markdown",
        fake_to_markdown,
    )

    result = document_parser.parse_document(path)

    assert result.extraction_method == "pymupdf4llm"
    assert result.page_count == 1
    assert "PDF parser test" in result.content