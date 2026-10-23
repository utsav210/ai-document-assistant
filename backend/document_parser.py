from __future__ import annotations

import re
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path
from threading import Lock

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}

_DOCling_CONVERTER = None
_DOCLING_INIT_LOCK = Lock()
_DOCLING_CONVERT_LOCK = Lock()


class DocumentParsingError(ValueError):
    """Raised when a document cannot be parsed reliably."""


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    source_name: str
    file_type: str
    content: str
    extraction_method: str
    page_count: int | None = None
    warnings: tuple[str, ...] = ()


def _normalize_markdown(content: str) -> str:
    content = content.replace("\x00", "")
    content = content.replace("\r\n", "\n").replace("\r", "\n")
    content = re.sub(r"\n{4,}", "\n\n\n", content)
    return content.strip()


def _get_docling_converter():
    """Initialize Docling once and reuse it for fallback conversions."""
    global _DOCling_CONVERTER

    if _DOCling_CONVERTER is None:
        with _DOCLING_INIT_LOCK:
            if _DOCling_CONVERTER is None:
                from docling.datamodel.base_models import InputFormat
                from docling.document_converter import DocumentConverter

                _DOCling_CONVERTER = DocumentConverter(
                    allowed_formats=[
                        InputFormat.PDF,
                        InputFormat.DOCX,
                    ]
                )

    return _DOCling_CONVERTER


def _docling_fallback(
    path: Path,
    file_type: str,
    reason: str,
    page_count: int | None = None,
) -> ParsedDocument:
    """Use Docling when the primary PDF/DOCX reader fails."""
    try:
        converter = _get_docling_converter()

        # Serialize fallback conversions when sharing one converter instance.
        with _DOCLING_CONVERT_LOCK:
            result = converter.convert(str(path))

        content = _normalize_markdown(result.document.export_to_markdown())

        if not content:
            raise DocumentParsingError("Docling returned empty document content.")

        return ParsedDocument(
            source_name=path.name,
            file_type=file_type,
            content=content,
            extraction_method="docling_fallback",
            page_count=page_count,
            warnings=(reason,),
        )

    except Exception as exc:
        raise DocumentParsingError(
            f"Could not parse '{path.name}'. "
            f"Primary extraction issue: {reason}. "
            f"Docling fallback failed: {exc}"
        ) from exc


def _pdf_page_chunks_to_markdown(chunks) -> str:
    """Combine PDF page chunks while preserving page markers."""
    if isinstance(chunks, str):
        return _normalize_markdown(chunks)

    sections = []

    for index, chunk in enumerate(chunks, start=1):
        page_text = chunk.get("text", "")
        page_text = _normalize_markdown(page_text)

        if not page_text:
            continue

        metadata = chunk.get("metadata") or {}
        page_number = metadata.get("page", index)

        sections.append(f"<!-- Page {page_number} -->\n\n{page_text}")

    return _normalize_markdown("\n\n".join(sections))


def _parse_pdf(path: Path) -> ParsedDocument:
    """Use PyMuPDF4LLM first, then Docling on failure or empty output."""
    try:
        import pymupdf4llm

        chunks = pymupdf4llm.to_markdown(
            str(path),
            page_chunks=True,
        )

        content = _pdf_page_chunks_to_markdown(chunks)

        if not content:
            raise DocumentParsingError("PyMuPDF4LLM returned empty content.")

        page_count = len(chunks) if isinstance(chunks, list) else None

        return ParsedDocument(
            source_name=path.name,
            file_type="pdf",
            content=content,
            extraction_method="pymupdf4llm",
            page_count=page_count,
        )

    except Exception as exc:  # noqa: BLE001
        return _docling_fallback(
            path=path,
            file_type="pdf",
            reason=f"PyMuPDF4LLM failed: {exc}",
        )


def _render_paragraph(paragraph) -> str:
    """Convert a DOCX paragraph to Markdown with basic inline formatting."""
    text_parts = []

    for run in paragraph.runs:
        text = run.text

        if not text:
            continue

        if run.bold:
            text = f"**{text}**"

        if run.italic:
            text = f"*{text}*"

        if run.underline:
            text = f"<u>{text}</u>"

        text_parts.append(text)

    text = "".join(text_parts).strip()

    if not text:
        return ""

    style_name = paragraph.style.name if paragraph.style is not None else ""

    heading_match = re.match(
        r"Heading\s*([1-6])",
        style_name,
        re.IGNORECASE,
    )

    if heading_match:
        level = int(heading_match.group(1))
        return f"{'#' * level} {text}"

    if style_name.lower() == "title":
        return f"# {text}"

    if "list bullet" in style_name.lower():
        return f"- {text}"

    if "list number" in style_name.lower():
        return f"1. {text}"

    return text


def _render_table(table) -> str:
    """Render a DOCX table as a Markdown table."""
    rows = []

    for row in table.rows:
        cells = []

        for cell in row.cells:
            value = cell.text.strip()
            value = value.replace("\n", "<br>")
            value = value.replace("|", r"\|")
            cells.append(value)

        rows.append(cells)

    if not rows:
        return ""

    column_count = max(len(row) for row in rows)

    if column_count == 0:
        return ""

    # Normalize irregular row widths.
    for row in rows:
        row.extend([""] * (column_count - len(row)))

    header = rows[0]
    separator = ["---"] * column_count

    output = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(separator) + " |",
    ]

    for row in rows[1:]:
        output.append("| " + " | ".join(row) + " |")

    return "\n".join(output)


def _parse_docx(path: Path) -> ParsedDocument:
    """Read DOCX body content in document order."""
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(str(path))
        sections = []

        for item in document.iter_inner_content():
            if isinstance(item, Paragraph):
                content = _render_paragraph(item)

            elif isinstance(item, Table):
                content = _render_table(item)

            else:
                continue

            if content.strip():
                sections.append(content.strip())

        markdown = _normalize_markdown("\n\n".join(sections))

        if not markdown:
            return _docling_fallback(
                path,
                "docx",
                "python-docx returned empty content.",
            )

        return ParsedDocument(
            source_name=path.name,
            file_type="docx",
            content=markdown,
            extraction_method="python-docx",
        )

    except Exception as exc:  # noqa: BLE001
        return _docling_fallback(
            path,
            "docx",
            f"python-docx failed: {exc}",
        )


def _read_text_file(path: Path) -> str:
    """Read text using common encodings while preserving line breaks."""
    raw = path.read_bytes()

    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16").replace("\r\n", "\n").replace("\r", "\n")

    if b"\x00" in raw[:512]:
        try:
            return raw.decode("utf-16").replace("\r\n", "\n").replace("\r", "\n")
        except UnicodeDecodeError:
            pass

    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding).replace("\r\n", "\n").replace("\r", "\n")
        except UnicodeDecodeError:
            continue

    raise DocumentParsingError(f"Could not decode text file: {path.name}")


def _parse_txt(path: Path) -> ParsedDocument:
    content = _read_text_file(path).strip()

    if not content:
        raise DocumentParsingError(f"Text file '{path.name}' is empty.")

    return ParsedDocument(
        source_name=path.name,
        file_type="txt",
        content=content,
        extraction_method="text_reader",
    )


def parse_document(file_path: str | Path) -> ParsedDocument:
    """Parse a supported document into a consistent Markdown-oriented result."""
    path = Path(file_path).expanduser().resolve()

    if not path.is_file():
        raise FileNotFoundError(f"Document not found: {path}")

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentParsingError(
            f"Unsupported file type '{extension}'. "
            f"Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if extension == ".pdf":
        return _parse_pdf(path)

    if extension == ".docx":
        return _parse_docx(path)

    return _parse_txt(path)


def lazy_parse_document(
    file_path: str | Path,
) -> Iterator[tuple[int | None, str]]:
    """Lazily yield (page_number, text) segments from a supported document."""
    path = Path(file_path).expanduser().resolve()

    if not path.is_file():
        raise FileNotFoundError(f"Document not found: {path}")

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentParsingError(
            f"Unsupported file type '{extension}'. "
            f"Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if extension == ".pdf":
        try:
            import pymupdf
            import pymupdf4llm

            doc = pymupdf.open(str(path))
            for page_idx in range(len(doc)):
                chunk = pymupdf4llm.to_markdown(doc, pages=[page_idx])
                normalized = _normalize_markdown(chunk)
                if normalized:
                    yield (page_idx + 1, normalized)
            return
        except Exception:  # noqa: BLE001
            parsed = _parse_pdf(path)
            matches = list(re.finditer(r"<!-- Page (\d+) -->", parsed.content))
            if not matches:
                yield (1, parsed.content)
                return
            for idx, match in enumerate(matches):
                end = (
                    matches[idx + 1].start()
                    if idx + 1 < len(matches)
                    else len(parsed.content)
                )
                text = _normalize_markdown(parsed.content[match.end() : end])
                if text:
                    yield (int(match.group(1)), text)
            return

    parsed = parse_document(path)
    yield (None, parsed.content)
